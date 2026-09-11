"""Validate completeness and consistency of the fictional TPRM project."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from project_data import ASSESSMENT, OVERALL_SCORE, REMEDIATION, RISKS, ROOT

REQUIRED_DIRS = [
    "00_Project_Overview",
    "01_Vendor_Profile",
    "02_Security_Questionnaire",
    "03_Vendor_Evidence",
    "04_Control_Assessment",
    "05_Risk_Assessment",
    "06_Risk_Register",
    "07_Remediation",
    "08_Risk_Treatment",
    "09_Vendor_Scorecard",
    "10_Executive_Report",
    "11_Final_Report",
    "12_Presentation",
    "13_Portfolio",
    "14_Scripts",
    "15_Final_Checks",
]

REQUIRED_FILES = [
    "README.md",
    "LICENSE",
    ".gitignore",
    "requirements.txt",
    "00_Project_Overview/Project_Charter.md",
    "00_Project_Overview/Business_Scenario.md",
    "00_Project_Overview/Assessment_Objectives.md",
    "00_Project_Overview/Scope_and_Assumptions.md",
    "00_Project_Overview/Methodology.md",
    "00_Project_Overview/Assessment_Timeline.md",
    "01_Vendor_Profile/Vendor_Profile.md",
    "01_Vendor_Profile/Vendor_Overview.md",
    "01_Vendor_Profile/Vendor_Data_Flow.md",
    "01_Vendor_Profile/Vendor_Architecture.md",
    "01_Vendor_Profile/Criticality_Assessment.md",
    "02_Security_Questionnaire/Vendor_Security_Questionnaire.xlsx",
    "02_Security_Questionnaire/Vendor_Response_Summary.xlsx",
    "02_Security_Questionnaire/Questionnaire_Methodology.md",
    "03_Vendor_Evidence/Information_Security_Policy.pdf",
    "03_Vendor_Evidence/Access_Control_Policy.pdf",
    "03_Vendor_Evidence/Incident_Response_Plan.pdf",
    "03_Vendor_Evidence/Business_Continuity_Plan.pdf",
    "03_Vendor_Evidence/Vulnerability_Management_Policy.pdf",
    "03_Vendor_Evidence/Data_Protection_Standard.pdf",
    "03_Vendor_Evidence/Penetration_Test_Summary.pdf",
    "03_Vendor_Evidence/SOC2_TypeI_Summary.pdf",
    "03_Vendor_Evidence/Backup_and_Recovery_Summary.pdf",
    "03_Vendor_Evidence/Subprocessor_Management_Summary.pdf",
    "04_Control_Assessment/Control_Assessment.xlsx",
    "04_Control_Assessment/Evidence_Mapping.xlsx",
    "04_Control_Assessment/Control_Assessment_Methodology.md",
    "05_Risk_Assessment/Risk_Scoring_Methodology.md",
    "05_Risk_Assessment/Risk_Matrix.md",
    "05_Risk_Assessment/Inherent_Risk_Assessment.xlsx",
    "05_Risk_Assessment/Residual_Risk_Assessment.xlsx",
    "05_Risk_Assessment/Risk_Assessment_Summary.md",
    "06_Risk_Register/Vendor_Risk_Register.xlsx",
    "06_Risk_Register/Risk_Register_Summary.md",
    "07_Remediation/Remediation_Plan.xlsx",
    "07_Remediation/Corrective_Action_Plan.md",
    "07_Remediation/Remediation_Prioritization.md",
    "08_Risk_Treatment/Risk_Treatment_Decisions.md",
    "08_Risk_Treatment/Risk_Acceptance_Record.md",
    "08_Risk_Treatment/Risk_Mitigation_Record.md",
    "08_Risk_Treatment/Residual_Risk_Acceptance.md",
    "09_Vendor_Scorecard/Vendor_Risk_Scorecard.xlsx",
    "09_Vendor_Scorecard/Vendor_Security_Scorecard.md",
    "09_Vendor_Scorecard/Scorecard_Explanation.md",
    "10_Executive_Report/Executive_Summary.md",
    "10_Executive_Report/Executive_Report.md",
    "10_Executive_Report/Management_Decision.md",
    "11_Final_Report/Third_Party_Vendor_Risk_Assessment_Report.md",
    "12_Presentation/Vendor_Risk_Assessment_Presentation.md",
    "13_Portfolio/Portfolio_Project_Summary.md",
    "13_Portfolio/CV_Project_Description.md",
    "13_Portfolio/LinkedIn_Project_Description.md",
    "13_Portfolio/Interview_Talking_Points.md",
    "14_Scripts/calculate_risk_scores.py",
    "14_Scripts/generate_scorecard.py",
    "14_Scripts/validate_project.py",
    "14_Scripts/README.md",
    "15_Final_Checks/Quality_Assurance_Checklist.md",
    "15_Final_Checks/Evidence_Traceability_Matrix.xlsx",
    "15_Final_Checks/Project_Completion_Checklist.md",
    "15_Final_Checks/Final_Findings.md",
]


def check(name, cond, failures):
    if cond:
        return "PASS"
    failures.append(name)
    return "FAIL"


def formulas_ok(path: Path) -> bool:
    wb = load_workbook(path)
    ws = wb["Risk_Register"]
    # Row 7 is first data row in generator
    for row in range(7, 15):
        inh = ws.cell(row, 9).value
        res = ws.cell(row, 13).value
        rate = ws.cell(row, 14).value
        if not (isinstance(inh, str) and inh.startswith("=")):
            return False
        if not (isinstance(res, str) and res.startswith("=")):
            return False
        if not (isinstance(rate, str) and rate.startswith("=")):
            return False
    return True


def main() -> int:
    failures: list[str] = []
    print("PROJECT VALIDATION")
    print("------------------")

    dirs_ok = all((ROOT / d).is_dir() for d in REQUIRED_DIRS)
    print(f"Directories: {check('directories', dirs_ok, failures)}")

    missing = [f for f in REQUIRED_FILES if not (ROOT / f).exists()]
    files_ok = not missing
    print(f"Core Documents: {check('core documents', files_ok, failures)}")
    if missing:
        for m in missing:
            print(f"  missing: {m}")

    ids = [r["id"] for r in RISKS]
    unique = len(ids) == len(set(ids))
    required_ids = {f"VR-00{i}" for i in range(1, 9)}
    has_all = required_ids <= set(ids)
    blanks = []
    for r in RISKS:
        for k in ("title", "description", "treatment", "evidence", "owner", "residual_rating"):
            if not str(r.get(k, "")).strip():
                blanks.append(f"{r['id']}.{k}")
    high = [r for r in RISKS if r["residual_rating"] == "High"]
    high_treated = all(r["treatment"] in {"Mitigate", "Accept", "Transfer", "Avoid"} for r in high)
    all_treated = all(r["treatment"] for r in RISKS)
    all_rated = all(r["residual_rating"] for r in RISKS)
    all_evidence = all(r["evidence"] for r in RISKS)
    rem_owners = all(c["owner"] and c["finding_id"] for c in REMEDIATION)
    rem_ids = {c["finding_id"] for c in REMEDIATION}
    risks_have_cap = all(r["id"] in rem_ids for r in RISKS)

    register = ROOT / "06_Risk_Register/Vendor_Risk_Register.xlsx"
    formula_ok = register.exists() and formulas_ok(register)

    report = (ROOT / "11_Final_Report/Third_Party_Vendor_Risk_Assessment_Report.md").read_text(encoding="utf-8") if (ROOT / "11_Final_Report/Third_Party_Vendor_Risk_Assessment_Report.md").exists() else ""
    report_has_risks = all(r["id"] in report for r in RISKS)
    report_decision = ASSESSMENT["recommendation"] in report and ASSESSMENT["overall_risk"] in report

    scorecard = ROOT / "09_Vendor_Scorecard/Vendor_Risk_Scorecard.xlsx"
    score_ok = scorecard.exists() and abs(OVERALL_SCORE - 74.4) < 0.05

    print(f"Risk Register: {check('risk register', unique and has_all and formula_ok and all_rated and all_treated, failures)}")
    print(f"Evidence Mapping: {check('evidence', all_evidence and (ROOT / '15_Final_Checks/Evidence_Traceability_Matrix.xlsx').exists(), failures)}")
    print(f"Remediation Mapping: {check('remediation', rem_owners and risks_have_cap, failures)}")
    print(f"Risk Treatment Mapping: {check('treatment', high_treated and all_treated, failures)}")
    print(f"Scorecard: {check('scorecard', score_ok, failures)}")
    extra_ok = (not blanks) and report_has_risks and report_decision
    if blanks:
        print("  blank fields: " + ", ".join(blanks))
        failures.append("blank fields")
    if not report_has_risks:
        print("  final report missing one or more risk IDs")
        failures.append("report risk coverage")
    if not report_decision:
        print("  final report missing recommendation or overall risk")
        failures.append("report decision")

    overall = "PASS" if not failures else "FAIL"
    print()
    print(f"OVERALL STATUS: {overall}")
    if overall != "PASS":
        print("Failed checks: " + ", ".join(failures))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
